"""
DOCX Report Assembly

Renders the technical due-diligence report from a report_content dict into a
.docx file: an Executive Summary followed by 10 numbered sections, each broken
into subsections with narrative prose, tables, and (for specific sections)
charts/diagrams. Every section is now a normal synthesized section - section 3
(Security & Compliance Findings) is fed by security_compliance_agent's native
osv-scanner/detect-secrets scan output, and section 5 (Recommended To-Be
Architecture) is always synthesized regardless of repo count - neither is a
hardcoded placeholder anymore. NOT_COVERED_SECTIONS/NOT_COVERED_NOTES are kept
as an empty-by-default mechanism for any future section that genuinely has no
v1 data source, rather than removed outright.

report_content["sections"]["<n>"] shape:
    {"subsections": [{"heading": "...", "narrative": "...", "table": {...}?}],
     "tables": [...]?, "chart": {...}?, "diagram": {...}?}
Chart/diagram PNGs are rendered on demand (tools/visuals.py) to
"<output_path stem>_images/" (sibling to the .docx); a section whose data is
missing or fails to render (e.g. Graphviz's `dot` binary isn't installed)
simply has no image for that slot - visuals.py's renderers never raise.
"""

import os
import logging
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from . import visuals

logger = logging.getLogger(__name__)
logger.setLevel(logging.INFO)

SECTION_TITLES = [
    "Current Architecture of the Codebase",
    "Business Logic & Domain Understanding",
    "Security & Compliance Findings",
    "Modernization Readiness",
    "Recommended To-Be Architecture",
    "Recommended AWS Service Usage",
    "Migration Roadmap",
    "Cost Benefit",
    "Performance & Reliability Benefit",
    "Risks & Mitigations",
]

# 1-indexed section numbers with no v1 data source. Empty now that security/
# compliance (3) and to-be architecture (5) both have real sources - kept as a
# mechanism, not deleted, for any future section that genuinely has none.
NOT_COVERED_SECTIONS = set()

NOT_COVERED_NOTES = {}

# Confirmed against an actual python-docx install this session - part of the default
# template's built-in table-style gallery. Fallback below is cheap insurance only.
TABLE_STYLE = "Light Grid Accent 1"


def _add_table(doc, headers, rows, style=TABLE_STYLE):
    if not headers or not rows:
        return
    table = doc.add_table(rows=1, cols=len(headers))
    try:
        table.style = style
    except KeyError:
        try:
            table.style = "Table Grid"
        except KeyError:
            pass  # no named table style available; render borderless rather than crash

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = str(h)
        for run in hdr_cells[i].paragraphs[0].runs:
            run.bold = True

    for row in rows:
        row_cells = table.add_row().cells
        for i in range(len(headers)):
            row_cells[i].text = str(row[i]) if i < len(row) else ""
    doc.add_paragraph()  # spacing after the table


def _add_image(doc, image_path, width_inches: float = 6.0):
    if not image_path or not os.path.isfile(image_path):
        return
    doc.add_picture(image_path, width=Inches(width_inches))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_narrative(doc, text):
    for para in str(text or "").split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())


def _render_section_visual(kind, data, images_dir, i):
    """kind: 'diagram' or 'chart'. Returns an image path or None."""
    if not data or not isinstance(data, dict):
        return None
    if kind == "diagram":
        return visuals.render_architecture_diagram(
            data, os.path.join(images_dir, f"section_{i}_diagram.png"))
    if kind == "chart":
        out = os.path.join(images_dir, f"section_{i}_chart.png")
        chart_type = data.get("type")
        if chart_type == "scorecard_bar":
            return visuals.render_scorecard_bar_chart(data, out)
        if chart_type == "roadmap_timeline":
            return visuals.render_roadmap_timeline_chart(data, out)
        logger.warning(f"Section {i}: unknown chart type {chart_type!r}, skipping")
        return None
    return None


def build_docx(report_content: dict, output_path: str) -> str:
    """
    Every section is now a normal synthesized section (see module docstring) -
    NOT_COVERED_SECTIONS is empty by default. Tolerates the old flat-string
    section shape and malformed section dicts defensively (falls back to
    plain-paragraph rendering) rather than crashing. Returns output_path.
    """
    doc = Document()
    images_dir = os.path.splitext(output_path)[0] + "_images"

    doc.add_heading('Technical Due Diligence Report', level=0)
    if report_content.get('client_name'):
        doc.add_paragraph(report_content['client_name'])
    if report_content.get('repos'):
        doc.add_paragraph('Repositories assessed: ' + ', '.join(report_content['repos']))

    if report_content.get('executive_summary'):
        doc.add_heading('Executive Summary', level=1)
        _add_narrative(doc, report_content['executive_summary'])

    sections = report_content.get('sections', {})

    for i, heading in enumerate(SECTION_TITLES, start=1):
        doc.add_heading(f"{i}. {heading}", level=1)
        if i in NOT_COVERED_SECTIONS:
            p = doc.add_paragraph()
            run = p.add_run(NOT_COVERED_NOTES[i])
            run.italic = True
            continue

        section = sections.get(str(i)) or sections.get(i)
        if not section:
            doc.add_paragraph("No content available for this section.")
            continue
        if isinstance(section, str):          # old flat-string shape, defensive
            section = {"subsections": [{"heading": "", "narrative": section}]}
        elif not isinstance(section, dict):
            doc.add_paragraph("No content available for this section.")
            continue

        subsections = section.get("subsections") or []
        if not subsections:
            doc.add_paragraph("No content available for this section.")

        for sub in subsections:
            if not isinstance(sub, dict):
                continue
            if sub.get("heading"):
                doc.add_heading(str(sub["heading"]), level=2)
            _add_narrative(doc, sub.get("narrative", ""))
            table = sub.get("table")
            if isinstance(table, dict) and table.get("headers") and table.get("rows"):
                _add_table(doc, table["headers"], table["rows"])

        for table in (section.get("tables") or []):
            if not isinstance(table, dict):
                continue
            if table.get("title"):
                doc.add_heading(str(table["title"]), level=3)
            if table.get("headers") and table.get("rows"):
                _add_table(doc, table["headers"], table["rows"])

        diagram = section.get("diagram")
        diagram_path = _render_section_visual("diagram", diagram, images_dir, i)
        if diagram_path:
            if isinstance(diagram, dict) and diagram.get("title"):
                doc.add_heading(str(diagram["title"]), level=3)
            _add_image(doc, diagram_path)

        chart_path = _render_section_visual("chart", section.get("chart"), images_dir, i)
        if chart_path:
            _add_image(doc, chart_path)

    if report_content.get('sources'):
        doc.add_heading('Sources', level=1)
        doc.add_paragraph(
            'AWS documentation referenced while preparing the recommendations in this report:'
        )
        for url in report_content['sources']:
            doc.add_paragraph(url, style='List Bullet')

    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
    doc.save(output_path)
    return output_path
